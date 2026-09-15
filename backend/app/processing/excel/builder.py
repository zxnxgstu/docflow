from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.db.models import Document


NAVY = "172033"
INDIGO = "5B5BD6"
PALE = "EEF2FF"
WHITE = "FFFFFF"


def _title(ws, title: str, subtitle: str):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(size=20, bold=True, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, color="64748B")


def _table(ws, headers: list[str], rows: list[list[object]], start_row: int = 4):
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(start_row, col, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    if rows:
        ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
        name = "DataTable" + str(abs(hash(ws.title)) % 100000)
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        ws.auto_filter.ref = ref
    ws.freeze_panes = f"A{start_row + 1}"
    for index, header in enumerate(headers, start=1):
        values = [str(header), *[str(row[index - 1] or "") for row in rows]]
        ws.column_dimensions[get_column_letter(index)].width = min(max(map(len, values)) + 3, 48)


def build_workbook(documents: list[Document], output: Path) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _title(summary, "DocFlow export", f"Generated {datetime.now():%Y-%m-%d %H:%M}")
    metrics = [
        ["Documents", len(documents)],
        ["Pages", sum(d.page_count for d in documents)],
        ["Extracted fields", sum(len(d.fields) for d in documents)],
        ["Extracted tables", sum(len(d.tables) for d in documents)],
        ["Edited values", sum(f.is_edited for d in documents for f in d.fields)],
    ]
    _table(summary, ["Metric", "Value"], metrics)

    docs_ws = wb.create_sheet("Documents")
    _title(docs_ws, "Documents", "Files included in this export")
    _table(docs_ws, ["Filename", "Status", "Pages", "Size (KB)", "Processed at"], [
        [d.original_filename, d.status, d.page_count, round(d.file_size / 1024, 1), d.processed_at.isoformat() if d.processed_at else ""]
        for d in documents
    ])

    fields_ws = wb.create_sheet("Extracted Fields")
    _title(fields_ws, "Extracted fields", "Validated key-value data")
    _table(fields_ws, ["Document", "Page", "Field", "Value", "Type", "Confidence", "Edited"], [
        [d.original_filename, f.page_number, f.field_name, f.field_value, f.field_type, f.confidence, "Yes" if f.is_edited else "No"]
        for d in documents for f in d.fields
    ])
    fields_ws.column_dimensions["D"].width = 40
    for cell in fields_ws["F"][4:]:
        cell.number_format = "0%"

    tables_ws = wb.create_sheet("Extracted Tables")
    _title(tables_ws, "Extracted tables", "Flattened rows from detected tables")
    table_rows = []
    for d in documents:
        for table in d.tables:
            for row_index, row in enumerate(table.data_json, start=1):
                table_rows.append([d.original_filename, table.page_number, table.table_index, row_index, " | ".join(map(str, row))])
    _table(tables_ws, ["Document", "Page", "Table", "Row", "Values"], table_rows)

    report = wb.create_sheet("Processing Report")
    _title(report, "Processing report", "Traceability and quality overview")
    _table(report, ["Document", "Status", "Fields", "Tables", "Average confidence", "Error"], [
        [d.original_filename, d.status, len(d.fields), len(d.tables),
         round(sum(f.confidence for f in d.fields) / len(d.fields), 2) if d.fields else 0,
         d.error_message or ""] for d in documents
    ])
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

